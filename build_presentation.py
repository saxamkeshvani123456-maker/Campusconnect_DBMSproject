from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.dml.color import RGBColor
from pathlib import Path

ROOT = Path(__file__).resolve().parent
OUT = ROOT / 'docs' / 'Project_Presentation.pptx'
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
NAVY = '202747'; NAVY2='292F52'; TEAL='8AD5C1'; TEALD='3C8E7D'; PAPER='F7F8F4'; WHITE='FFFFFF'; INK='252943'; MUTED='777B90'; PALE='EAF5F0'; LAV='EBE8FF'; ORANGE='FFD4A8'; LINE='E0E3E8'; RED='D77C70'


def rgb(h): return RGBColor.from_string(h)

def rect(slide,x,y,w,h,fill, radius=False, line=None, transparency=0):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid(); shp.fill.fore_color.rgb=rgb(fill); shp.fill.transparency=transparency
    if line:
        shp.line.color.rgb=rgb(line); shp.line.width=Pt(1)
    else: shp.line.fill.background()
    if radius:
        try: shp.adjustments[0] = 0.15
        except Exception: pass
    return shp

def text(slide, x,y,w,h, value, size=16, color=INK, bold=False, font='Aptos', align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.MIDDLE, margin=0, italic=False):
    shape=slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf=shape.text_frame; tf.clear(); tf.word_wrap=True
    tf.margin_left=Inches(margin); tf.margin_right=Inches(margin); tf.margin_top=Inches(0.02); tf.margin_bottom=Inches(0.02); tf.vertical_anchor=valign
    p=tf.paragraphs[0]; p.alignment=align
    r=p.add_run(); r.text=str(value); r.font.name=font; r.font.size=Pt(size); r.font.bold=bold; r.font.italic=italic; r.font.color.rgb=rgb(color)
    return shape

def add_bg(slide, color=PAPER):
    bg=slide.background.fill; bg.solid(); bg.fore_color.rgb=rgb(color)

def label(slide,x,y,w,s, color=TEALD, dark=False):
    text(slide,x,y,w,.23,s.upper(),9,color,True,font='Aptos')

def title(slide, overline, heading, sub=None, dark=False):
    c = WHITE if dark else NAVY
    m = 'B9C1D7' if dark else MUTED
    label(slide,.62,.38,5,overline, TEAL if dark else TEALD, dark)
    text(slide,.62,.68,12,.56,heading,27,c,True,font='Aptos Display')
    if sub: text(slide,.64,1.25,12,.35,sub,11,m)

def footer(slide,num,dark=False):
    c='AEB5CC' if dark else '9498A8'
    text(slide,.63,7.17,8,.18,'CAMPUSCONNECT  /  CS203 DBMS PROJECT',7.5,c,True)
    text(slide,12.1,7.13,.55,.22,f'{num:02d}  /  09',7.5,c,True,align=PP_ALIGN.RIGHT)

def pill(slide,x,y,w,h,value,fill=PALE,color=TEALD,fontsize=9):
    rect(slide,x,y,w,h,fill,True)
    text(slide,x+.08,y,w-.16,h,value,fontsize,color,True,align=PP_ALIGN.CENTER)

def card(slide,x,y,w,h,head,body,accent=TEAL,fill=WHITE, icon=None):
    rect(slide,x,y,w,h,fill,True,line=LINE)
    rect(slide,x,y,.06,h,accent,True)
    left=x+.22
    if icon:
        rect(slide,left,y+.22,.42,.42,accent,True)
        text(slide,left,y+.22,.42,.42,icon,15,NAVY,True,align=PP_ALIGN.CENTER)
        left += .58
    text(slide,left,y+.2,w-(left-x)-.16,.42,head,14,NAVY,True,font='Aptos Display')
    text(slide,x+.23,y+.73,w-.45,h-.87,body,10,MUTED,valign=MSO_ANCHOR.TOP)

# 1. Cover
s=prs.slides.add_slide(BLANK); add_bg(s,NAVY)
rect(s,8.35,0,4.98,7.05,NAVY2,True)
# abstract orbit motif
for x,y,w,h,c in [(8.7,1.0,3.8,3.8,'334064'),(9.15,1.45,2.9,2.9,'3C486D'),(9.65,1.95,1.9,1.9,'8AD5C1')]:
    shp=s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(w), Inches(h)); shp.fill.solid(); shp.fill.fore_color.rgb=rgb(c); shp.line.fill.background()
rect(s,9.35,2.0,2.75,2.3,WHITE,True)
text(s,9.62,2.35,2.2,.6,'cc.',38,NAVY,True,font='Aptos Display')
text(s,9.62,3.15,2.15,.45,'CAMPUS / CONNECTED',9,TEALD,True)
text(s,.72,.62,4,.35,'SVNIT SURAT  ·  CSE',11,TEAL,True)
text(s,.72,1.55,7.1,1.55,'CampusConnect',44,WHITE,True,font='Aptos Display')
text(s,.75,3.08,6.6,.75,'A student campus events & community portal',20,'D9DEED',False,font='Aptos Display')
rect(s,.75,4.25,.8,.05,TEAL,True)
text(s,.75,4.58,6.2,.75,'DATABASE MANAGEMENT SYSTEMS  ·  CS203',12,'C3C9DB',True)
text(s,.75,5.25,6.3,.35,'B.Tech. II CSE  |  Semester III  |  Academic Year 2026–27',10,'AEB6CC')
text(s,.75,6.15,6.2,.42,'Team: [registered team name]  •  Members: [add names & roll numbers]',9,'AEB6CC')
footer(s,1,True)

# 2. Problem and response
s=prs.slides.add_slide(BLANK); add_bg(s)
title(s,'01 / PROJECT REQUIREMENTS','Campus life is connected. Campus information often is not.','CampusConnect puts the everyday student experience into one database-backed portal.')
card(s,.7,1.9,3.85,2.3,'Scattered information','Events and notices can be spread across chats, posters, club pages, and informal networks.',ORANGE,WHITE,'!')
card(s,4.75,1.9,3.85,2.3,'Hard to discover','Students may not know which clubs are active or where to find a relevant campus activity.',LAV,WHITE,'⌕')
card(s,8.8,1.9,3.85,2.3,'Manual coordination','Organisers need a consistent way to publish events, maintain registrations, and share updates.',TEAL,WHITE,'↗')
rect(s,.7,4.65,11.95,1.55,NAVY,True)
text(s,1.05,4.9,2.4,.28,'THE PROPOSED RESPONSE',9,TEAL,True)
text(s,1.05,5.25,7.95,.55,'One student-friendly home for events, clubs, registrations, and announcements.',19,WHITE,True,font='Aptos Display')
pill(s,10.05,5.13,2.1,.45,'DESIGNED FOR CS203',fill='314064',color=WHITE,fontsize=9)
footer(s,2)

# 3 Requirements/features
s=prs.slides.add_slide(BLANK); add_bg(s)
title(s,'02 / SCOPE','Four core experiences. One shared data model.','The prototype focuses on explainable database operations rather than an oversized feature list.')
card(s,.7,1.85,2.88,2.0,'Discover events','Search future and past events; view date, venue, host club, and capacity.',TEAL,WHITE,'✳')
card(s,3.77,1.85,2.88,2.0,'Join communities','Browse clubs and create or end a membership relationship.',LAV,WHITE,'◎')
card(s,6.84,1.85,2.88,2.0,'Stay in the loop','Read announcements based on audience and optional expiry.',ORANGE,WHITE,'↗')
card(s,9.91,1.85,2.72,2.0,'Manage the portal','Admin can publish, edit, or delete event and campus content.',TEALD,WHITE,'⌘')
text(s,.72,4.45,3.1,.3,'KEY REQUIREMENTS',10,TEALD,True)
for i,(head,body) in enumerate([
    ('Account & profile','Registration, login, profile updates'),
    ('Capacity-aware registration','No duplicate pair; check available spots'),
    ('Role-based administration','Admin routes protected on server'),
    ('Integrity by design','PK, FK, CHECK, UNIQUE, indexes, views')]):
    x=.72+(i%2)*6.08; y=4.88+(i//2)*.82
    rect(s,x,y,5.65,.62,WHITE,True,line=LINE)
    text(s,x+.18,y+.09,2.15,.23,head,10,NAVY,True)
    text(s,x+2.25,y+.08,3.2,.34,body,9,MUTED)
footer(s,3)

# 4 Roles + journey
s=prs.slides.add_slide(BLANK); add_bg(s)
title(s,'03 / USER EXPERIENCE','Different roles. Clear boundaries.','A student-facing experience, with a separate administrator control room.')
rect(s,.72,1.83,5.7,4.7,WHITE,True,line=LINE)
pill(s,1.02,2.12,1.3,.34,'STUDENT',fill=PALE,color=TEALD)
text(s,1.02,2.62,4.9,.43,'Find what is happening.',20,NAVY,True,font='Aptos Display')
for idx, (num,heading_,desc) in enumerate([
    ('01','Create account','Name, email, department, study year'),
    ('02','Explore events & clubs','Search events, join clubs, read notices'),
    ('03','Take part','Register, cancel, update personal profile'),
]):
    y=3.35+idx*.82
    rect(s,1.02,y,.42,.42,PALE,True); text(s,1.02,y,.42,.42,num,9,TEALD,True,align=PP_ALIGN.CENTER)
    text(s,1.62,y-.01,3.8,.25,heading_,11,NAVY,True)
    text(s,1.62,y+.27,4.25,.24,desc,9,MUTED)
rect(s,6.72,1.83,5.9,4.7,NAVY,True)
pill(s,7.03,2.12,1.72,.34,'ADMINISTRATOR',fill='354061',color=TEAL,fontsize=8.5)
text(s,7.03,2.62,4.9,.43,'Keep the campus feed fresh.',20,WHITE,True,font='Aptos Display')
for idx, (num,heading_,desc) in enumerate([
    ('01','Publish','Create events, clubs, and announcements'),
    ('02','Maintain','Edit event details and remove stale content'),
    ('03','Review data','See event counts and registration totals'),
]):
    y=3.35+idx*.82
    rect(s,7.03,y,.42,.42,'364164',True); text(s,7.03,y,.42,.42,num,9,TEAL,True,align=PP_ALIGN.CENTER)
    text(s,7.63,y-.01,3.8,.25,heading_,11,WHITE,True)
    text(s,7.63,y+.27,4.25,.24,desc,9,'C1C8D9')
footer(s,4)

# 5 ER diagram
s=prs.slides.add_slide(BLANK); add_bg(s,WHITE)
label(s,.58,.24,4,'04 / CONCEPTUAL DESIGN',TEALD)
text(s,.58,.48,12,.4,'The ER model',24,NAVY,True,font='Aptos Display')
text(s,.6,.9,12,.25,'Six relations cover identity, activities, and the two many-to-many relationships.',10,MUTED)
erd=ROOT/'docs'/'ER_Diagram.png'
s.shapes.add_picture(str(erd), Inches(1.1), Inches(1.25), width=Inches(11.05))
text(s,.65,6.96,8,.2,'PK = primary key  ·  FK = foreign key  ·  Relationship cardinalities are also listed in the report.',8,MUTED)
footer(s,5)

# 6 Relational schema normalization
s=prs.slides.add_slide(BLANK); add_bg(s)
title(s,'05 / LOGICAL DESIGN','The relationship tables do the heavy lifting.','The logical schema maps each entity to a table and resolves many-to-many relationships explicitly.')
# table schema visual blocks
for x,y,w,head,body,col in [
    (.72,1.85,5.7,'USERS','user_id (PK) · email (UNIQUE)\nfull_name · password_hash · department\nstudy_year · role · created_at',TEAL),
    (6.82,1.85,5.8,'CLUBS / EVENTS','clubs: club_id (PK), name (UNIQUE)\nevents: event_id (PK), event_date, capacity\nclub_id / created_by (FK)',LAV),
    (.72,3.55,5.7,'CLUB_MEMBERSHIPS','Composite PK: (user_id, club_id)\njoined_at · status\nResolves users ↔ clubs (M:N)',ORANGE),
    (6.82,3.55,5.8,'EVENT_REGISTRATIONS','Composite PK: (event_id, user_id)\nregistered_at · status\nResolves users ↔ events (M:N)',TEALD),
]:
    rect(s,x,y,w,1.32,WHITE,True,line=LINE); rect(s,x,y,w,.38,col,True)
    text(s,x+.18,y+.04,w-.36,.28,head,11,NAVY,True)
    text(s,x+.2,y+.48,w-.4,.73,body,10,MUTED,valign=MSO_ANCHOR.TOP)
rect(s,.72,5.18,11.9,1.08,NAVY,True)
text(s,1.0,5.36,1.5,.25,'NORMALISATION',9,TEAL,True)
text(s,1.0,5.68,10.7,.3,'1NF: atomic fields   ·   2NF: relationship attributes depend on full composite keys   ·   3NF: club/event/profile facts stored once',10,WHITE,True)
text(s,.77,6.46,11.5,.3,'Announcements are stored separately; event registration counts are derived in a view instead of duplicated in the Events table.',9,MUTED,italic=True)
footer(s,6)

# 7 physical design
s=prs.slides.add_slide(BLANK); add_bg(s)
title(s,'06 / PHYSICAL DESIGN','Integrity rules are part of the database—not just the interface.','The package includes SQLite for a quick local demo and MySQL 8 DDL for SQL demonstration.')
card(s,.72,1.9,3.78,2.05,'Keys & relationships','Primary keys, foreign keys, composite keys, and cascades prevent orphan or duplicate relationship rows.',TEAL,WHITE,'⌘')
card(s,4.77,1.9,3.78,2.05,'Domain constraints','UNIQUE email/name, CHECK ranges, and valid role/status/audience values restrict invalid data.',ORANGE,WHITE,'✓')
card(s,8.82,1.9,3.78,2.05,'Query performance','Indexes support event date filtering, membership counts, per-user history, and recent notices.',LAV,WHITE,'↗')
rect(s,.72,4.35,7.25,1.8,NAVY,True)
text(s,1.0,4.6,2.5,.25,'REPORTING VIEWS',9,TEAL,True)
text(s,1.0,4.98,6.45,.35,'upcoming_events_view',15,WHITE,True,font='Aptos Display')
text(s,1.0,5.38,6.45,.3,'Future events + host club + active registration count',9,'C4CBDB')
rect(s,8.23,4.35,4.37,1.8,PALE,True)
text(s,8.52,4.6,3.8,.25,'SQL ARTEFACTS',9,TEALD,True)
for i,line in enumerate(['schema_mysql.sql','schema_sqlite.sql','crud_examples.sql']):
    text(s,8.52,4.95+i*.32,3.7,.25,'•  '+line,10,NAVY,True)
footer(s,7)

# 8 Prototype architecture and demo
s=prs.slides.add_slide(BLANK); add_bg(s)
title(s,'07 / PROTOTYPE','From the browser to the database, in one simple flow.','A Flask web application connects the interface to SQL tables and business rules.')
# 3-level architecture
for x,y,w,h,head,body,col in [
    (.85,1.95,3.25,1.35,'01  BROWSER','HTML templates · CSS · JavaScript',LAV),
    (5.05,1.95,3.25,1.35,'02  FLASK ROUTES','Session login · validation · role checks',PALE),
    (9.25,1.95,3.25,1.35,'03  SQL DATABASE','Foreign keys · constraints · indexes · views',ORANGE),
]:
    rect(s,x,y,w,h,WHITE,True,line=LINE); rect(s,x,y,.08,h,col,True)
    text(s,x+.22,y+.2,w-.4,.32,head,13,NAVY,True)
    text(s,x+.22,y+.64,w-.45,.48,body,10,MUTED,valign=MSO_ANCHOR.TOP)
for x in [4.28,8.48]:
    text(s,x,2.25,.58,.5,'→',24,TEALD,True,align=PP_ALIGN.CENTER)
text(s,.87,3.75,5,.3,'LIVE DEMO SEQUENCE',10,TEALD,True)
steps=[('A','Student account','Sign up and log in'),('B','Participation','Register for an event + join a club'),('C','Admin control','Create / edit event + publish notice'),('D','SQL proof','Show stored rows and reporting views')]
for i,(letter,head,desc) in enumerate(steps):
    x=.85+i*3.1
    rect(s,x,4.25,2.82,1.35,WHITE,True,line=LINE)
    pill(s,x+.16,4.42,.38,.33,letter,fill=PALE,color=TEALD)
    text(s,x+.68,4.4,1.95,.31,head,11,NAVY,True)
    text(s,x+.17,4.91,2.48,.48,desc,9,MUTED,valign=MSO_ANCHOR.TOP)
rect(s,.85,5.95,11.22,.48,'FFF3E8',True)
text(s,1.08,6.05,10.75,.24,'Demo credentials are for local use only. Run the browser workflow on the presentation machine before evaluation.',9,'946D4A',True)
footer(s,8)

# 9 tests and close
s=prs.slides.add_slide(BLANK); add_bg(s,NAVY)
label(s,.68,.46,4,'08 / TESTING & WRAP-UP',TEAL,True)
text(s,.68,.82,10.8,.58,'Built around the requirements.',28,WHITE,True,font='Aptos Display')
text(s,.7,1.48,11,.35,'The database schema checks pass; the team should complete browser workflow checks locally before presenting.',11,'C5CADD')
# left verification card
rect(s,.72,2.2,5.55,3.6,NAVY2,True)
text(s,1.02,2.48,4.8,.34,'VERIFIED AT SCHEMA LEVEL',10,TEAL,True)
for i,line in enumerate(['Six expected tables and two views created','Composite key rejects duplicate registrations','Foreign keys block orphan memberships','CHECK constraint rejects capacity = 0','Status update and reporting count tests pass']):
    y=3.0+i*.48
    text(s,1.02,y,.25,.25,'✓',13,TEAL,True)
    text(s,1.38,y,4.5,.3,line,10,WHITE)
# right action card
rect(s,6.65,2.2,5.95,3.6,WHITE,True)
text(s,7.0,2.48,4.9,.34,'BEFORE FINAL SUBMISSION',10,TEALD,True)
for i,line in enumerate(['Fill in team name, names, and roll numbers','Run Flask app and complete browser checks','Open MySQL schema in the agreed SQL tool','Capture screenshots and update test status','Rehearse the live demo and explain the ER model']):
    y=3.0+i*.48
    text(s,7.0,y,.25,.25,'□',12,TEALD,True)
    text(s,7.36,y,4.7,.3,line,10,INK)
text(s,.72,6.25,9.5,.48,'CampusConnect  ·  One campus, more possibilities.',17,WHITE,True,font='Aptos Display')
text(s,10.3,6.34,2.3,.28,'THANK YOU  ✦',10,TEAL,True,align=PP_ALIGN.RIGHT)
footer(s,9,True)

prs.core_properties.title='CampusConnect - DBMS Project Presentation'
prs.core_properties.subject='CS203 Database Management Systems Project, SVNIT Surat, 2026-27'
prs.core_properties.author='CampusConnect Project Team'
prs.core_properties.keywords='CampusConnect, DBMS, ER diagram, SQL, SVNIT Surat'
prs.save(OUT)
print(OUT)
